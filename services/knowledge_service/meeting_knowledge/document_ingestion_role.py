"""Register and read supported meeting-document sources."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from services.knowledge_service.store import KnowledgeApplicationDataStore
from shared.contracts import (
    LocatorType,
    SourceDocument,
    SourceDocumentStatus,
    SourceType,
)


SUPPORTED_SOURCE_TYPES = {
    ".pdf": SourceType.PDF,
    ".pptx": SourceType.PPTX,
}


class DocumentReadError(RuntimeError):
    """Raised when a supported source cannot be parsed safely."""


@dataclass(frozen=True)
class ExtractedBlock:
    locator_type: LocatorType
    locator_value: str
    order: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    blocks: list[ExtractedBlock]
    warnings: list[str] = field(default_factory=list)

    @property
    def character_count(self) -> int:
        return sum(len(block.text) for block in self.blocks)


@dataclass(frozen=True)
class RegistrationResult:
    document: SourceDocument
    duplicate: bool


def calculate_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def register_source_document(
    path: Path,
    store: KnowledgeApplicationDataStore,
) -> RegistrationResult:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SOURCE_TYPES:
        raise ValueError(f"unsupported meeting document type: {suffix or '<none>'}")
    if not path.is_file():
        raise FileNotFoundError(f"meeting document does not exist: {path}")

    content_hash = calculate_sha256(path)
    existing = store.get_source_document_by_hash(content_hash)
    if existing:
        if existing.status is SourceDocumentStatus.PROCESSED:
            return RegistrationResult(document=existing, duplicate=True)
        retryable = existing.model_copy(
            update={
                "source_path": str(path),
                "status": SourceDocumentStatus.REGISTERED,
                "error_message": None,
            }
        )
        store.save_source_document(retryable)
        return RegistrationResult(document=retryable, duplicate=False)

    document = SourceDocument(
        source_id=uuid5(NAMESPACE_URL, f"meeting-source:{content_hash}"),
        source_type=SUPPORTED_SOURCE_TYPES[suffix],
        title=path.stem,
        source_path=str(path),
        content_hash=content_hash,
        status=SourceDocumentStatus.REGISTERED,
    )
    store.save_source_document(document)
    return RegistrationResult(document=document, duplicate=False)


def read_meeting_document(
    path: Path,
    document: SourceDocument,
) -> ExtractedDocument:
    try:
        if calculate_sha256(path) != document.content_hash:
            raise ValueError("source content changed after registration")
        if document.source_type is SourceType.PDF:
            return _read_pdf(path)
        if document.source_type is SourceType.PPTX:
            return _read_pptx(path)
    except Exception as exc:
        raise DocumentReadError(f"cannot read {path.name}: {exc}") from exc
    raise DocumentReadError(f"unsupported source type: {document.source_type}")


def _read_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise ValueError("encrypted PDF requires a password")

    blocks: list[ExtractedBlock] = []
    warnings: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text(extraction_mode="layout") or ""
        if not text.strip():
            warnings.append(f"PDF page {page_number} contains no extractable text")
            continue
        blocks.append(
            ExtractedBlock(
                locator_type=LocatorType.PDF_PAGE,
                locator_value=str(page_number),
                order=page_number,
                text=text,
            )
        )
    if not blocks:
        raise ValueError("PDF contains no extractable text")
    return ExtractedDocument(blocks=blocks, warnings=warnings)


def _shape_text(shape: object) -> str:
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        child_shapes = sorted(
            shape.shapes,
            key=lambda child: (int(child.top), int(child.left)),
        )
        return "\n".join(filter(None, (_shape_text(child) for child in child_shapes)))

    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        return "\n".join(rows)

    if getattr(shape, "has_text_frame", False):
        paragraphs = []
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if not text:
                text = paragraph.text.strip()
            if text:
                paragraphs.append(text)
        return "\n".join(paragraphs)

    return ""


def _read_pptx(path: Path) -> ExtractedDocument:
    presentation = Presentation(str(path))
    blocks: list[ExtractedBlock] = []
    warnings: list[str] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        ordered_shapes = sorted(
            slide.shapes,
            key=lambda shape: (int(shape.top), int(shape.left)),
        )
        slide_block_count = 0
        for shape_order, shape in enumerate(ordered_shapes, start=1):
            text = _shape_text(shape).strip()
            if not text:
                continue
            slide_block_count += 1
            blocks.append(
                ExtractedBlock(
                    locator_type=LocatorType.PPTX_SLIDE,
                    locator_value=str(slide_number),
                    order=(slide_number * 10_000) + shape_order,
                    text=text,
                )
            )
        if slide_block_count == 0:
            warnings.append(f"PPTX slide {slide_number} contains no extractable text")

    if not blocks:
        raise ValueError("PPTX contains no extractable text")
    return ExtractedDocument(blocks=blocks, warnings=warnings)
